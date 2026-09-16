# -*- coding: utf-8 -*-
# 82_md_to_docx_v2.py - 参照第二版格式：Unicode 公式 + 4 张图 + 作者块
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"D:\paper\astra version\00_full_draft.md"
OUT = r"D:\paper\astra version\manuscript_astra_version.docx"
FIGDIR = r"D:\paper\02_analysis\figures"

def latex_to_unicode(s):
    greek = [(r'\alpha','α'),(r'\beta','β'),(r'\gamma','γ'),(r'\delta','δ'),
        (r'\epsilon','ε'),(r'\zeta','ζ'),(r'\eta','η'),(r'\theta','θ'),
        (r'\kappa','κ'),(r'\lambda','λ'),(r'\mu','μ'),(r'\pi','π'),
        (r'\sigma','σ'),(r'\tau','τ'),(r'\phi','φ'),(r'\chi','χ'),
        (r'\psi','ψ'),(r'\omega','ω'),
        (r'\Gamma','Γ'),(r'\Delta','Δ'),(r'\Theta','Θ'),(r'\Lambda','Λ'),
        (r'\Pi','Π'),(r'\Sigma','Σ'),(r'\Phi','Φ'),(r'\Psi','Ψ'),(r'\Omega','Ω')]
    ops = [(r'\times','×'),(r'\cdot','·'),(r'\approx','≈'),(r'\leq','≤'),
        (r'\geq','≥'),(r'\iff','⇔'),(r'\pm','±'),(r'\infty','∞'),
        (r'\equiv','≡'),(r'\neq','≠'),(r'\sim','∼'),(r'\wedge','∧'),
        (r'\propto','∝'),(r'\partial','∂'),(r'\sum','Σ'),(r'\prod','Π'),
        (r'\sqrt','√'),(r'\rightarrow','→'),(r'\to','→')]
    for a,b in greek+ops:
        s = s.replace(a,b)
    s = re.sub(r'\\text\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\mathbf\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\mathcal\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\mathrm\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\frac\{([^}]*)\}\{([^}]*)\}', r'\1/\2', s)
    s = re.sub(r'\\sqrt\{([^}]*)\}', r'√\1', s)
    s = re.sub(r'\\hat\{([^}]*)\}', lambda m: m.group(1) + chr(0x0302), s)
    s = re.sub(r'\\bar\{([^}]*)\}', lambda m: m.group(1) + chr(0x0304), s)
    s = re.sub(r'\\tilde\{([^}]*)\}', lambda m: m.group(1) + chr(0x0303), s)
    s = re.sub(r'\\overset\{[^}]*\}\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\xrightarrow\{[^}]*\}', '→', s)
    s = re.sub(r'\\in\b', '∈', s)
    s = re.sub(r'\\le\b', '≤', s)
    s = re.sub(r'\\ge\b', '≥', s)
    for cmd in ['min','max','ln','log','exp','Var','sign','Delta']:
        s = s.replace('\\'+cmd, cmd)
    s = re.sub(r'_\{([^}]*)\}', r'_\1', s)
    s = re.sub(r'\^\{([^}]*)\}', r'^\1', s)
    s = re.sub(r'\\([a-zA-Z]+)', r'\1', s)
    s = s.replace('$', '')
    return s

with open(SRC, encoding="utf-8") as f:
    md = f.read()

doc = Document()

# ===== 作者块 =====
title_text = "Censoring-geometry collapse of mutation\u00d7lineage interaction estimands in pharmacogenomic panels: bidirectional bias and an operational estimator-selection rule"
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(title_text); r.bold = True; r.font.size = Pt(14)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Jun Bao\u00b9, Chao Song\u00b9 and Zhongze Gu\u00b9,*").font.size = Pt(12)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\u00b9 School of Biological Science and Medical Engineering, Southeast University, Nanjing 210096, China").font.size = Pt(10)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("* Correspondence: gu@seu.edu.cn (Z.G.)").font.size = Pt(10)
doc.add_paragraph()

# Graphical Abstract
try:
    doc.add_picture(FIGDIR + r"\graphical_abstract.png", width=Inches(6.2))
    last_para = doc.paragraphs[-1]; last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
except Exception as e:
    print("GA 图失败:", e)

FIGMAP = {
    "(Figure 2)": (FIGDIR + r"\fig2_sigma_bias.png", 5.5),
    "(Figure 1)": (FIGDIR + r"\fig1_censoring_geometry.png", 5.5),
    "(Figure 3)": (FIGDIR + r"\fig3_consort.png", 5.5),
}

def add_runs(paragraph, text):
    text = latex_to_unicode(text)
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part:
            paragraph.add_run(part)

lines = md.split("\n")
i = 0
first_h1 = True
while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped.startswith("# ") and first_h1:
        first_h1 = False; i += 1; continue

    # 表格
    if stripped.startswith("|") and i+1 < len(lines) and lines[i+1].strip().startswith("|"):
        tbl_lines = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            tbl_lines.append(lines[i].strip()); i += 1
        def parse_row(r): return [c.strip() for c in r.strip("|").split("|")]
        header = parse_row(tbl_lines[0])
        data_rows = [parse_row(r) for r in tbl_lines[2:] if not re.match(r'^\|[\s:\-|]+\|$', r)]
        table = doc.add_table(rows=1+len(data_rows), cols=len(header))
        table.style = 'Light Grid Accent 1'
        for j,h in enumerate(header):
            table.rows[0].cells[j].paragraphs[0].add_run(latex_to_unicode(h.replace('**',''))).bold = True
        for ri,row in enumerate(data_rows, start=1):
            for j,ct in enumerate(row):
                if j < len(header):
                    table.rows[ri].cells[j].paragraphs[0].add_run(latex_to_unicode(ct.replace('**','')))
        doc.add_paragraph()
        continue

    if stripped.startswith("### "):
        doc.add_heading(latex_to_unicode(stripped[4:]), level=3)
    elif stripped.startswith("## "):
        doc.add_heading(latex_to_unicode(stripped[3:]), level=2)
    elif stripped.startswith("# "):
        doc.add_heading(latex_to_unicode(stripped[2:]), level=1)
    elif stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(latex_to_unicode(stripped[2:-2]))
    elif stripped.startswith("- "):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
        add_runs(p, "\u2022 " + stripped[2:])
    elif stripped == "":
        pass
    else:
        p = doc.add_paragraph()
        add_runs(p, stripped)
        # 图片插入：段落含 (Figure N) 标记时，紧跟插图
        for marker, (fpath, w) in FIGMAP.items():
            if marker in stripped:
                try:
                    pic = doc.add_picture(fpath, width=Inches(w))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    print(f"插图失败 {marker}:", e)

    i += 1

doc.save(OUT)
import os
print(f"已生成: {OUT}")
print(f"文件大小: {os.path.getsize(OUT)} bytes")
