# -*- coding: utf-8 -*-
# 81_md_to_docx.py - astra version full_draft.md -> Word（含作者信息）
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"D:\paper\astra version\00_full_draft.md"
OUT = r"D:\paper\astra version\manuscript_astra_version.docx"

with open(SRC, encoding="utf-8") as f:
    md = f.read()

doc = Document()

# ===== 标题块 =====
title_text = "Censoring-geometry collapse of mutation\u00d7lineage interaction estimands in pharmacogenomic panels: bidirectional bias and an operational estimator-selection rule"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(title_text)
r.bold = True
r.font.size = Pt(14)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Jun Bao\u00b9, Chao Song\u00b9 and Zhongze Gu\u00b9,*").font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\u00b9 School of Biological Science and Medical Engineering, Southeast University, Nanjing 210096, China")
r.font.size = Pt(10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("* Correspondence: gu@seu.edu.cn (Z.G.)")
r.font.size = Pt(10)

doc.add_paragraph()

# ===== 解析 body =====
def add_runs(paragraph, text):
    """解析 **bold** 和 $...$ 行内公式"""
    text = re.sub(r'\$([^$]+)\$', r'\1', text)  # 去 $ 保留 LaTeX
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)

lines = md.split("\n")
i = 0
first_h1 = True
while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    # 跳过第一个 H1（标题，已入标题块）
    if stripped.startswith("# ") and first_h1:
        first_h1 = False
        i += 1
        continue

    # 表格
    if stripped.startswith("|") and i + 1 < len(lines) and lines[i+1].strip().startswith("|"):
        # 收集表格行
        tbl_lines = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            tbl_lines.append(lines[i].strip())
            i += 1
        # 解析：第一行表头，第二行分隔（跳过），其余数据
        def parse_row(r):
            cells = [c.strip() for c in r.strip("|").split("|")]
            return cells
        header = parse_row(tbl_lines[0])
        data_rows = [parse_row(r) for r in tbl_lines[2:] if not re.match(r'^\|[\s:\-|]+\|$', r)]
        table = doc.add_table(rows=1 + len(data_rows), cols=len(header))
        table.style = 'Light Grid Accent 1'
        for j, h in enumerate(header):
            cell = table.rows[0].cells[j]
            cell.paragraphs[0].add_run(re.sub(r'\*\*', '', h)).bold = True
        for ri, row in enumerate(data_rows, start=1):
            for j, cell_text in enumerate(row):
                if j < len(header):
                    table.rows[ri].cells[j].paragraphs[0].add_run(re.sub(r'\*\*', '', cell_text))
        doc.add_paragraph()
        continue

    # 标题
    if stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=3)
    elif stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=2)
    elif stripped.startswith("# "):
        doc.add_heading(stripped[2:], level=1)
    # 块公式 $$...$$
    elif stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(stripped[2:-2])
    # 列表
    elif stripped.startswith("- "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        add_runs(p, "\u2022 " + stripped[2:])
    # 空行
    elif stripped == "":
        pass
    # 普通段落
    else:
        p = doc.add_paragraph()
        add_runs(p, stripped)

    i += 1

doc.save(OUT)
import os
print(f"已生成: {OUT}")
print(f"文件大小: {os.path.getsize(OUT)} bytes")
