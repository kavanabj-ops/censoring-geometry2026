# -*- coding: utf-8 -*-
# 84_cover_letter.py - BMC Bioinformatics cover letter
from docx import Document
from docx.shared import Pt

OUT = r"D:\paper\astra version\cover_letter_BMC_Bioinformatics.docx"

doc = Document()

# 正文（信件格式）
paras = [
    ("Dear Editor,", False),
    ("", False),
    ("We are pleased to submit our manuscript entitled \u201cWhen the wild type hits the ceiling: identifiability of mutation\u2013lineage interactions in pharmacogenomic panels\u201d for consideration as a Methodology article in BMC Bioinformatics.", False),
    ("", False),
    ("This manuscript addresses a practical problem in pharmacogenomic data analysis. When a targeted drug's wild-type cell lines are insensitive to the drug, their IC50 values are recorded at the assay ceiling (right-censored), so 94\u201395% of them carry no quantitative information. We show that this censoring geometry\u2014not the analyst's choice of estimator\u2014is what makes the mutation \u00d7 lineage interaction effect practically unidentifiable, and that the two standard workarounds (extrapolating the ceiling, or capping at it) bias the result in opposite, predictable directions.", False),
    ("", False),
    ("The contribution is threefold and directly relevant to the readership of BMC Bioinformatics: (i) an analytic characterization of the identifiability boundary, including a sign-reversal condition and a minimum-detectable-effect; (ii) a data-driven calibration of the key extrapolation bound using GDSC's own censoring distribution; and (iii) an operational two-step rule\u2014check per-cell sample size, then censoring\u2014that a pharmacogenomics practitioner can apply immediately. We complement the theory with simulations and a census of GDSC2/GDSC1 showing a three-level bottleneck (94.5% without a hotspot mutation, 45% sample-size-insufficient, 36.4% non-informative by censoring).", False),
    ("", False),
    ("All analyses use publicly available data (GDSC1/GDSC2, CCLE/DepMap), and the code reproduces every figure and table. We are committed to BMC's reproducibility standards and will deposit the code in a public repository with a DOI upon acceptance.", False),
    ("", False),
    ("This work is original, has not been published elsewhere, and is not under consideration by another journal. All authors have approved the manuscript and agree to its submission. We declare no competing interests. The study uses only publicly available cell-line data and did not involve human or animal subjects.", False),
    ("", False),
    ("Thank you for considering our manuscript. We look forward to your response.", False),
    ("", False),
    ("Sincerely,", False),
    ("Zhongze Gu, Ph.D.", True),
    ("Professor, School of Biological Science and Medical Engineering, Southeast University", False),
    ("Corresponding author; gu@seu.edu.cn", False),
]

for text, bold in paras:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    if bold:
        run.bold = True

doc.save(OUT)
import os
print(f"已生成: {OUT} ({os.path.getsize(OUT)} bytes)")
